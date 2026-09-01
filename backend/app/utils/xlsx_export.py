"""Shared helper to build a simple, styled .xlsx workbook for data exports.

Mirrors the openpyxl convention already used by the reports module (bold header row,
frozen header, reasonable column widths) so all Excel exports look consistent.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Sequence


def build_xlsx(sheet_title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> BytesIO:
    """Return an in-memory .xlsx (BytesIO, positioned at 0) with one styled sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Sheet1")[:31]  # Excel sheet-name limit

    ws.append(list(headers))
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = header_align

    for row in rows:
        ws.append(list(row))

    ws.freeze_panes = "A2"

    # Size each column to the longest value (bounded), for readability.
    for col_idx in range(1, len(headers) + 1):
        longest = len(str(headers[col_idx - 1]))
        for row in rows:
            value = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            longest = max(longest, len(str(value)) if value is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(10, longest + 2), 50)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
